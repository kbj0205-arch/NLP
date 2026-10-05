"""CT2S-1.0 인간 주석지(셀 안 줄바꿈 형식)의 기계 점검.

사용: python tools/check_human_annotation.py 주석지.xlsx [--sheet 주석1_001-50 ...]

판정 내용은 보지 않는다. 채점·병합 때 깨질 수 있는 형식만 점검한다.
- 단위원문·술어군·명시주어·주제부·규칙 열의 줄 수가 같은가(빈 줄·끝 줄바꿈 포함)
- 명시주어·주제부·단위원문 조각이 원문에 글자 그대로 있는가(띄어쓰기 차이는 따로 알림)
- 주제부가 그 단위의 시작부터 이어지는가(T2·T3·T4)
- 검토상태를 골랐는가
"""
import argparse
import re
import sys

from openpyxl import load_workbook

EMPTY_VALUES = {"없음", "명시 없음", "유보", "-", ""}
COLS = {"id": 1, "text": 2, "unit": 6, "pred": 7, "subj": 8, "theme": 9, "rule": 10, "status": 11}
LINE_COLS = ["unit", "pred", "subj", "theme", "rule"]
LABEL = {"unit": "단위원문", "pred": "술어군", "subj": "명시주어", "theme": "주제부", "rule": "규칙"}


def lines(value):
    if value is None:
        return None
    return str(value).replace("\r\n", "\n").split("\n")


def squash(s):
    return re.sub(r"\s+", "", s)


def locate(fragment, text):
    """원문에서 fragment를 찾는다. (상태, 원문 쪽 문자열)"""
    frag = fragment.strip().rstrip(".")
    if not frag:
        return "ok", frag
    if frag in text:
        return "ok", frag
    # 띄어쓰기만 다른지: 공백을 뺀 문자열로 대응 위치를 찾아 원문 구간을 돌려준다
    target = squash(frag)
    idx_map = [i for i, ch in enumerate(text) if not ch.isspace()]
    flat = "".join(text[i] for i in idx_map)
    pos = flat.find(target)
    if pos >= 0:
        start, end = idx_map[pos], idx_map[pos + len(target) - 1] + 1
        return "spacing", text[start:end]
    return "missing", None


def check_row(text, cells):
    issues = []
    split = {k: lines(cells[k]) for k in LINE_COLS}
    present = {k: v for k, v in split.items() if v is not None}
    if not present:
        return issues

    for k, v in present.items():
        if len(v) > 1 and v[-1].strip() == "":
            issues.append(f"{LABEL[k]}: 끝에 빈 줄이 있다(줄 수가 하나 늘어남)")
        if any(x.strip() == "" for x in v[:-1]):
            issues.append(f"{LABEL[k]}: 중간에 빈 줄이 있다")
    counts = {LABEL[k]: len([x for x in v if x.strip() != ""]) for k, v in present.items()}
    if len(set(counts.values())) > 1:
        issues.append("줄 수가 다르다: " + ", ".join(f"{k} {n}" for k, n in counts.items()))

    def nonblank(k):
        return [x for x in (split[k] or []) if x.strip() != ""]

    units, subjs, themes, rules = nonblank("unit"), nonblank("subj"), nonblank("theme"), nonblank("rule")

    for i, u in enumerate(units, 1):
        for piece in u.split("||"):
            st, orig = locate(piece, text)
            if st == "spacing":
                issues.append(f"U{i} 단위원문: 띄어쓰기가 원문과 다르다 → 원문 표기 '{orig}'")
            elif st == "missing":
                issues.append(f"U{i} 단위원문: 원문에 그대로 없다('{piece.strip()}'). 떨어진 조각이면 '||'로 나눈다")

    for label, vals in (("명시주어", subjs), ("주제부", themes)):
        for i, v in enumerate(vals, 1):
            if v.strip() in EMPTY_VALUES:
                continue
            st, orig = locate(v, text)
            if st == "spacing":
                issues.append(f"U{i} {label}: 띄어쓰기가 원문과 다르다 → 원문 표기 '{orig}'")
            elif st == "missing":
                issues.append(f"U{i} {label}: 원문에 이어진 문자열로 없다('{v.strip()}'). 빠진 말이 없는지 확인")

    for i, (u, t) in enumerate(zip(units, themes), 1):
        r = rules[i - 1] if i - 1 < len(rules) else ""
        if t.strip() in EMPTY_VALUES or not re.search(r"T[234]", r):
            continue
        first = squash(u.split("||")[0])
        if not first.startswith(squash(t).rstrip(".")):
            issues.append(f"U{i} 주제부: 단위의 시작부터 이어지지 않는다(T2–T4는 절의 시작부터)")

    if not cells["status"]:
        issues.append("검토상태가 비어 있다(확정·부분확정·유보 중 선택)")
    return issues


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", nargs="*")
    args = ap.parse_args()

    wb = load_workbook(args.xlsx, data_only=True)
    sheets = [wb[s] for s in args.sheet] if args.sheet else [ws for ws in wb.worksheets if ws.title.startswith("주석")]
    n_rows = n_bad = 0
    for ws in sheets:
        for r in range(2, ws.max_row + 1):
            cells = {k: ws.cell(r, c).value for k, c in COLS.items()}
            if not any(cells[k] for k in LINE_COLS):
                continue
            n_rows += 1
            issues = check_row(str(cells["text"] or ""), cells)
            if issues:
                n_bad += 1
                print(f"[{ws.title} {r}행 {cells['id']}]")
                for msg in issues:
                    print(f"  - {msg}")
    print(f"\n점검한 문장 {n_rows}개, 고칠 곳이 있는 문장 {n_bad}개")


if __name__ == "__main__":
    sys.exit(main())
